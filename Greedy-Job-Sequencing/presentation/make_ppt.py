from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Color Palette
COLOR_BG = RGBColor(248, 250, 252)        # Light background
COLOR_PRIMARY = RGBColor(15, 23, 42)     # Dark slate headers
COLOR_ACCENT = RGBColor(37, 99, 235)     # Royal blue accent
COLOR_MUTED = RGBColor(100, 116, 139)    # Muted gray text
COLOR_CARD = RGBColor(255, 255, 255)     # Card fill
COLOR_BORDER = RGBColor(226, 232, 240)   # Light gray border

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_BG

def add_header(slide, title_text, category_text="CS ALGORITHMS"):
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11), Inches(1.0))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = category_text.upper()
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT
    p0.font.name = 'Arial'
    
    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.size = Pt(26)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_PRIMARY
    p1.font.name = 'Arial'

# ==========================================
# SLIDE 1: Title
# ==========================================
slide1 = prs.slides.add_slide(blank_layout)
set_slide_background(slide1)

# Title box
tx_title = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(2.5))
tf1 = tx_title.text_frame
tf1.word_wrap = True

p_main = tf1.paragraphs[0]
p_main.text = "Greedy Job Sequencing\nwith Deadlines & Profits"
p_main.font.size = Pt(40)
p_main.font.bold = True
p_main.font.color.rgb = COLOR_PRIMARY
p_main.font.name = 'Arial'

# Subtitle / Info
tx_roll = slide1.shapes.add_textbox(Inches(1.0), Inches(5.0), Inches(6.0), Inches(1.5))
tf_roll = tx_roll.text_frame
p_roll = tf_roll.paragraphs[0]
p_roll.text = "Roll Number: 24WH1A0530"
p_roll.font.size = Pt(18)
p_roll.font.color.rgb = COLOR_MUTED
p_roll.font.name = 'Arial'

# ==========================================
# SLIDE 2: Problem Statement
# ==========================================
slide2 = prs.slides.add_slide(blank_layout)
set_slide_background(slide2)
add_header(slide2, "Problem Statement & Key Definitions")

items = [
    ("Job (J_i)", "An independent task requiring 1 unit of execution time on a single resource."),
    ("Deadline (d_i)", "The latest time slot by which a job must complete (1-indexed)."),
    ("Profit (p_i)", "The monetary reward earned strictly upon completing the job within its deadline."),
    ("Objective Goal", "Maximize total profit by choosing a subset of jobs that execute without slot overlaps.")
]

for idx, (term, desc) in enumerate(items):
    col = idx % 2
    row = idx // 2
    x = Inches(0.8 + col * 5.8)
    y = Inches(1.8 + row * 2.5)
    
    shape = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(5.4), Inches(2.1))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_CARD
    shape.line.color.rgb = COLOR_BORDER
    
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = term
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(14)
    p2.font.color.rgb = COLOR_PRIMARY

# ==========================================
# SLIDE 3: Greedy Approach
# ==========================================
slide3 = prs.slides.add_slide(blank_layout)
set_slide_background(slide3)
add_header(slide3, "The Greedy Decision Strategy")

steps = [
    ("1. Sort Jobs", "Order all given jobs in descending order based on profit."),
    ("2. Select Job", "Iterate from highest to lowest profit, evaluating each job."),
    ("3. Find Slot", "Search for the latest available slot t ≤ deadline."),
    ("4. Schedule", "Assign job to slot if available; skip if all slots t ≤ d are taken.")
]

for idx, (title, desc) in enumerate(steps):
    x = Inches(0.8 + idx * 2.9)
    y = Inches(2.5)
    
    shape = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.6), Inches(3.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_CARD
    shape.line.color.rgb = COLOR_ACCENT if idx == 0 else COLOR_BORDER
    
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLOR_MUTED

# ==========================================
# SLIDE 4: Example Dataset
# ==========================================
slide4 = prs.slides.add_slide(blank_layout)
set_slide_background(slide4)
add_header(slide4, "Example Dataset & Sorted Order")

table_shape = slide4.shapes.add_table(6, 4, Inches(2.0), Inches(2.0), Inches(9.333), Inches(4.0))
table = table_shape.table

headers = ["Job Identifier", "Deadline (d_i)", "Profit (p_i)", "Greedy Priority Order"]
data = [
    ["J1", "2", "100", "Rank 1 (Highest Profit)"],
    ["J3", "2", "27",  "Rank 2"],
    ["J4", "1", "25",  "Rank 3"],
    ["J2", "1", "19",  "Rank 4"],
    ["J5", "3", "15",  "Rank 5 (Lowest Profit)"]
]

for col_idx, header in enumerate(headers):
    cell = table.cell(0, col_idx)
    cell.text = header
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_PRIMARY
    for p in cell.text_frame.paragraphs:
        p.font.bold = True
        p.font.color.rgb = COLOR_CARD
        p.font.size = Pt(14)

for row_idx, row_data in enumerate(data):
    for col_idx, val in enumerate(row_data):
        cell = table.cell(row_idx + 1, col_idx)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CARD if row_idx % 2 == 0 else COLOR_BG
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(13)
            p.font.color.rgb = COLOR_PRIMARY

# ==========================================
# SLIDE 5: Step-by-Step Execution
# ==========================================
slide5 = prs.slides.add_slide(blank_layout)
set_slide_background(slide5)
add_header(slide5, "Step-by-Step Scheduling Trace")

trace_steps = [
    ("Initial State", "[ __ ]  [ __ ]  [ __ ]", "Max Deadline = 3 slots"),
    ("Consider J1 (p=100, d=2)", "[ __ ]  [ J1 ]  [ __ ]", "Assigned to latest open slot (Slot 2)"),
    ("Consider J3 (p=27, d=2)",  "[ J3 ]  [ J1 ]  [ __ ]", "Slot 2 taken -> Assigned to Slot 1"),
    ("Consider J4 (p=25, d=1)",  "[ J3 ]  [ J1 ]  [ __ ]", "Slot 1 taken -> Cannot schedule"),
    ("Consider J2 (p=19, d=1)",  "[ J3 ]  [ J1 ]  [ __ ]", "Slot 1 taken -> Cannot schedule"),
    ("Consider J5 (p=15, d=3)",  "[ J3 ]  [ J1 ]  [ J5 ]", "Assigned to Slot 3")
]

for idx, (title, state, note) in enumerate(trace_steps):
    y = Inches(1.8 + idx * 0.8)
    
    shape = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), y, Inches(11.733), Inches(0.7))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_CARD
    shape.line.color.rgb = COLOR_BORDER
    
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"{title}   ➔   "
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_PRIMARY
    
    p_state = tf.add_paragraph()
    p_state.text = f"Timeline: {state}    ({note})"
    p_state.font.size = Pt(12)
    p_state.font.color.rgb = COLOR_ACCENT

# Result Callout
tx_res = slide5.shapes.add_textbox(Inches(0.8), Inches(6.6), Inches(11.733), Inches(0.6))
tf_res = tx_res.text_frame
p_res = tf_res.paragraphs[0]
p_res.text = "Scheduled Jobs: J3, J1, J5   |   Maximum Profit = 100 + 27 + 15 = 142"
p_res.font.bold = True
p_res.font.size = Pt(16)
p_res.font.color.rgb = COLOR_PRIMARY

# ==========================================
# SLIDE 6: Flowchart
# ==========================================
slide6 = prs.slides.add_slide(blank_layout)
set_slide_background(slide6)
add_header(slide6, "Algorithm Flowchart")

flow_nodes = [
    ("Start", "Input Jobs(p, d)"),
    ("Sort", "Sort Jobs by Profit ↓"),
    ("Loop", "For each Job J_i in Sorted Order"),
    ("Check", "Find latest free slot t ≤ d_i"),
    ("Decision", "Is Slot Free?"),
    ("Assign / Skip", "YES: Assign J_i to Slot\nNO: Skip Job"),
    ("End", "Return Scheduled Schedule & Total Profit")
]

for idx, (stage, desc) in enumerate(flow_nodes):
    x = Inches(0.8 + idx * 1.65)
    y = Inches(3.2)
    
    shape = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(1.5), Inches(1.8))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_CARD
    shape.line.color.rgb = COLOR_ACCENT
    
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = stage
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_ACCENT
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(10)
    p2.font.color.rgb = COLOR_PRIMARY

# ==========================================
# SLIDE 7: Algorithm & Complexity
# ==========================================
slide7 = prs.slides.add_slide(blank_layout)
set_slide_background(slide7)
add_header(slide7, "Pseudocode & Complexity Analysis")

# Pseudocode Box
shape_code = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(6.5), Inches(5.0))
shape_code.fill.solid()
shape_code.fill.fore_color.rgb = RGBColor(15, 23, 42)
tf_code = shape_code.text_frame
tf_code.word_wrap = True

code_text = """JobSequencing(Jobs, n):
    Sort Jobs in decreasing order of Profit
    max_d = max(d_i for all jobs)
    slot = array of size (max_d + 1), initialized to FREE
    
    for i = 1 to n:
        for j = min(max_d, Jobs[i].deadline) down to 1:
            if slot[j] is FREE:
                slot[j] = Jobs[i].id
                total_profit += Jobs[i].profit
                break"""

p_c = tf_code.paragraphs[0]
p_c.text = code_text
p_c.font.size = Pt(12)
p_c.font.name = 'Courier New'
p_c.font.color.rgb = RGBColor(241, 245, 249)

# Complexity Box
shape_comp = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.6), Inches(1.8), Inches(4.9), Inches(5.0))
shape_comp.fill.solid()
shape_comp.fill.fore_color.rgb = COLOR_CARD
shape_comp.line.color.rgb = COLOR_BORDER
tf_comp = shape_comp.text_frame
tf_comp.word_wrap = True

p_h = tf_comp.paragraphs[0]
p_h.text = "Complexity Breakdown\n"
p_h.font.bold = True
p_h.font.size = Pt(18)
p_h.font.color.rgb = COLOR_PRIMARY

bullets = [
    "Time Complexity:",
    "• Sorting: O(n log n)",
    "• Slot Searching: O(n × d)",
    "• Overall Time: O(n log n + n × d)",
    "• (Can be optimized to O(n log n) using Disjoint Set Union / DSU)\n",
    "Space Complexity:",
    "• Slot Array: O(d) where d = max deadline"
]

for b in bullets:
    p = tf_comp.add_paragraph()
    p.text = b
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_PRIMARY if ":" in b else COLOR_MUTED
    if ":" in b:
        p.font.bold = True

# Save output
prs.save("Job_Sequencing_Presentation.pptx")
print("Presentation generated successfully: Job_Sequencing_Presentation.pptx")